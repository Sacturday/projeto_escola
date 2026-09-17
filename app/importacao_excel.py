from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
import hashlib
import re
import unicodedata

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from zipfile import BadZipFile
import xlrd
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from . import models

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 50_000
MAX_COLUMNS = 64
MAX_ERRORS_RETURNED = 100
ALLOWED_EXTENSIONS = {"xls", "xlsx"}

ALIASES = {
    "NOME_DA_ESCOLA": "ESCOLA_NOME", "NOME_ESCOLA": "ESCOLA_NOME",
    "NOME_DA_TURMA": "TURMA_NOME", "NOME_TURMA": "TURMA_NOME",
    "ANO": "ANO_LETIVO", "MATRICULA_ESCOLAR": "MATRICULA", "NOME_ALUNO": "NOME",
    "DATA_NASC": "DATA_NASCIMENTO", "DATA_DE_NASCIMENTO": "DATA_NASCIMENTO",
    "INICIO": "DATA_INICIO", "FIM": "DATA_FIM",
}

TIPOS_COLUNAS = {
    models.TipoImportacao.ESCOLAS: {"NOME": True, "CODIGO": True},
    models.TipoImportacao.TURMAS: {"ESCOLA_CODIGO": True, "NOME": True, "ANO_LETIVO": True, "TURNO": False},
    models.TipoImportacao.ALUNOS: {"ESCOLA_CODIGO": True, "MATRICULA": True, "NOME": True, "DATA_NASCIMENTO": False},
    models.TipoImportacao.MATRICULAS: {
        "ESCOLA_CODIGO": True, "MATRICULA": True, "TURMA_NOME": True, "ANO_LETIVO": True,
        "DATA_INICIO": False, "DATA_FIM": False,
    },
}


def normalizar_cabecalho(valor: object) -> str:
    texto = unicodedata.normalize("NFKD", str(valor or ""))
    texto = "".join(ch for ch in texto if not unicodedata.combining(ch))
    texto = re.sub(r"[^A-Za-z0-9]+", "_", texto.upper()).strip("_")
    return ALIASES.get(texto, texto)


def normalizar_valor(valor: object) -> object:
    return valor.strip() if isinstance(valor, str) else valor


def parse_data(valor: object) -> date | None:
    if valor in (None, ""):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, (int, float)):
        base = date(1899, 12, 30)
        try:
            return base.fromordinal(base.toordinal() + int(valor))
        except (OverflowError, ValueError):
            raise ValueError("data serial do Excel inválida")
    texto = str(valor).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"data inválida: {texto}")


def parse_int(valor: object, campo: str) -> int:
    try:
        if isinstance(valor, float) and valor.is_integer():
            return int(valor)
        return int(str(valor).strip())
    except (TypeError, ValueError):
        raise ValueError(f"{campo} deve ser inteiro")


async def _ler_bytes(arquivo: UploadFile) -> tuple[bytes, str, str]:
    nome = arquivo.filename or "arquivo"
    extensao = nome.lower().rsplit(".", 1)[-1] if "." in nome else ""
    if extensao not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=400, detail="Envie um arquivo .xls ou .xlsx")
    dados = await arquivo.read()
    if len(dados) > MAX_FILE_BYTES:
        raise HTTPException(status_code=413, detail="Arquivo excede o limite de 10 MB")
    if not dados:
        raise HTTPException(status_code=400, detail="Arquivo vazio")
    return dados, extensao, hashlib.sha256(dados).hexdigest()


async def ler_primeira_planilha(arquivo: UploadFile) -> tuple[str, list[dict[str, object]], str]:
    dados, extensao, sha256 = await _ler_bytes(arquivo)
    headers: list[str] = []
    rows: list[dict[str, object]] = []

    if extensao == "xlsx":
        try:
            wb = load_workbook(BytesIO(dados), read_only=True, data_only=True)
        except (BadZipFile, ValueError, OSError) as exc:
            raise HTTPException(status_code=400, detail=f"XLSX inválido ou corrompido: {exc}") from exc
        try:
            ws = next((sheet for sheet in wb.worksheets if sheet.max_row and sheet.max_column), None)
            if ws is None:
                raise HTTPException(status_code=400, detail="Nenhuma planilha preenchida foi encontrada")
            if ws.max_column > MAX_COLUMNS:
                raise HTTPException(status_code=413, detail=f"Planilha excede o limite de {MAX_COLUMNS} colunas")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i == 0:
                    headers = [normalizar_cabecalho(v) for v in row]
                    duplicadas = sorted({h for h in headers if h and headers.count(h) > 1})
                    if duplicadas:
                        raise HTTPException(status_code=422, detail="Cabeçalhos duplicados: " + ", ".join(duplicadas))
                    continue
                if i > MAX_ROWS:
                    raise HTTPException(status_code=413, detail=f"Planilha excede o limite de {MAX_ROWS} linhas")
                valores = list(row[:len(headers)])
                if not any(v not in (None, "") for v in valores):
                    continue
                rows.append({h: normalizar_valor(valores[idx] if idx < len(valores) else None) for idx, h in enumerate(headers) if h})
            return ws.title, rows, sha256
        finally:
            wb.close()

    try:
        wb = xlrd.open_workbook(file_contents=dados, on_demand=True)
    except (xlrd.XLRDError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=f"XLS inválido ou corrompido: {exc}") from exc
    try:
        sheet = next((s for s in wb.sheets() if s.nrows and s.ncols), None)
        if sheet is None:
            raise HTTPException(status_code=400, detail="Nenhuma planilha preenchida foi encontrada")
        if sheet.ncols > MAX_COLUMNS:
            raise HTTPException(status_code=413, detail=f"Planilha excede o limite de {MAX_COLUMNS} colunas")
        headers = [normalizar_cabecalho(sheet.cell_value(0, c)) for c in range(sheet.ncols)]
        duplicadas = sorted({h for h in headers if h and headers.count(h) > 1})
        if duplicadas:
            raise HTTPException(status_code=422, detail="Cabeçalhos duplicados: " + ", ".join(duplicadas))
        for r in range(1, sheet.nrows):
            if r > MAX_ROWS:
                raise HTTPException(status_code=413, detail=f"Planilha excede o limite de {MAX_ROWS} linhas")
            if not any(sheet.cell_value(r, c) not in (None, "") for c in range(sheet.ncols)):
                continue
            registro: dict[str, object] = {}
            for c, h in enumerate(headers):
                if not h:
                    continue
                cell = sheet.cell(r, c)
                valor: object = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    valor = xlrd.xldate_as_datetime(cell.value, wb.datemode)
                registro[h] = normalizar_valor(valor)
            rows.append(registro)
        return sheet.name, rows, sha256
    finally:
        wb.release_resources()


def validar_colunas(tipo: models.TipoImportacao, rows: list[dict[str, object]]) -> list[str]:
    if not rows:
        return ["A planilha não possui linhas de dados."]
    headers = list(rows[0].keys())
    duplicadas = sorted({h for h in headers if headers.count(h) > 1})
    if duplicadas:
        return ["Cabeçalhos duplicados: " + ", ".join(duplicadas)]
    obrigatorias = {c for c, obrigatoria in TIPOS_COLUNAS[tipo].items() if obrigatoria}
    faltantes = sorted(obrigatorias - set(headers))
    if faltantes:
        return ["Colunas obrigatórias ausentes: " + ", ".join(faltantes)]
    return []


async def validar_dados(tipo: models.TipoImportacao, rows: list[dict[str, object]], db: AsyncSession) -> list[dict[str, object]]:
    erros: list[dict[str, object]] = []
    def add(idx: int, erro: str):
        if len(erros) < MAX_ERRORS_RETURNED:
            erros.append({"linha": idx, "erro": erro})

    escolas = {e.codigo for e in (await db.execute(select(models.Escola.codigo))).all() if e[0]}
    turmas = {(t.escola_id, t.nome, t.ano_letivo) for t in (await db.execute(select(models.Turma.escola_id, models.Turma.nome, models.Turma.ano_letivo))).all()}
    alunos = {(a.escola_id, a.matricula) for a in (await db.execute(select(models.Aluno.escola_id, models.Aluno.matricula))).all()}

    vistos: set[tuple] = set()
    for idx, row in enumerate(rows, start=2):
        try:
            if tipo == models.TipoImportacao.ESCOLAS:
                nome = str(row.get("NOME") or "").strip(); codigo = str(row.get("CODIGO") or "").strip()
                if not nome: raise ValueError("NOME é obrigatório")
                if not codigo: raise ValueError("CODIGO é obrigatório")
                if codigo in vistos: raise ValueError(f"CODIGO duplicado na planilha: {codigo}")
                vistos.add(codigo)
            elif tipo == models.TipoImportacao.TURMAS:
                codigo = str(row.get("ESCOLA_CODIGO") or "").strip(); nome = str(row.get("NOME") or "").strip(); ano = parse_int(row.get("ANO_LETIVO"), "ANO_LETIVO")
                if not codigo: raise ValueError("ESCOLA_CODIGO é obrigatório")
                if not nome: raise ValueError("NOME é obrigatório")
                if codigo not in escolas and codigo not in {str(r.get("ESCOLA_CODIGO") or "").strip() for r in rows if tipo == models.TipoImportacao.TURMAS}: raise ValueError(f"escola não encontrada: {codigo}")
                chave=(codigo,nome,ano)
                if chave in vistos: raise ValueError("turma duplicada na planilha")
                vistos.add(chave)
            elif tipo == models.TipoImportacao.ALUNOS:
                codigo = str(row.get("ESCOLA_CODIGO") or "").strip(); matricula = str(row.get("MATRICULA") or "").strip(); nome = str(row.get("NOME") or "").strip()
                if not codigo: raise ValueError("ESCOLA_CODIGO é obrigatório")
                if not matricula: raise ValueError("MATRICULA é obrigatória")
                if not nome: raise ValueError("NOME é obrigatório")
                if row.get("DATA_NASCIMENTO") not in (None, ""): parse_data(row.get("DATA_NASCIMENTO"))
                chave=(codigo,matricula)
                if chave in vistos: raise ValueError("matrícula duplicada na planilha")
                vistos.add(chave)
            elif tipo == models.TipoImportacao.MATRICULAS:
                codigo = str(row.get("ESCOLA_CODIGO") or "").strip(); matricula = str(row.get("MATRICULA") or "").strip(); turma_nome = str(row.get("TURMA_NOME") or "").strip(); ano = parse_int(row.get("ANO_LETIVO"), "ANO_LETIVO")
                for campo, val in (("ESCOLA_CODIGO", codigo), ("MATRICULA", matricula), ("TURMA_NOME", turma_nome)):
                    if not val: raise ValueError(f"{campo} é obrigatório")
                inicio = parse_data(row.get("DATA_INICIO")) if row.get("DATA_INICIO") not in (None, "") else date(ano,1,1)
                fim = parse_data(row.get("DATA_FIM")) if row.get("DATA_FIM") not in (None, "") else None
                if fim and fim < inicio: raise ValueError("DATA_FIM não pode ser anterior a DATA_INICIO")
                if codigo not in escolas: raise ValueError(f"escola não encontrada: {codigo}")
        except (ValueError, TypeError) as exc:
            add(idx, str(exc))

    # Validação relacional mais forte após a primeira passada; permite importar escolas
    # e, separadamente, as entidades dependentes usando o código.
    if tipo in (models.TipoImportacao.TURMAS, models.TipoImportacao.ALUNOS, models.TipoImportacao.MATRICULAS):
        for idx, row in enumerate(rows, start=2):
            if len(erros) >= MAX_ERRORS_RETURNED: break
            codigo = str(row.get("ESCOLA_CODIGO") or "").strip()
            if codigo and codigo not in escolas:
                add(idx, f"escola não encontrada no banco: {codigo}")
            if tipo == models.TipoImportacao.MATRICULAS and codigo:
                matricula = str(row.get("MATRICULA") or "").strip()
                aluno = (await db.execute(select(models.Aluno.id).join(models.Escola).where(models.Escola.codigo==codigo, models.Aluno.matricula==matricula))).scalar_one_or_none()
                if aluno is None:
                    add(idx, f"aluno não encontrado no banco: {matricula}")
                ano = parse_int(row.get("ANO_LETIVO"), "ANO_LETIVO")
                nome_turma = str(row.get("TURMA_NOME") or "").strip()
                turma = (await db.execute(select(models.Turma.id).join(models.Escola).where(models.Escola.codigo==codigo, models.Turma.nome==nome_turma, models.Turma.ano_letivo==ano))).scalar_one_or_none()
                if turma is None:
                    add(idx, f"turma não encontrada no banco: {nome_turma}/{ano}")
    return erros


async def importar(tipo: models.TipoImportacao, rows: list[dict[str, object]], db: AsyncSession) -> tuple[int, list[dict[str, object]]]:
    erros: list[dict[str, object]] = []
    importados = 0
    for idx, row in enumerate(rows, start=2):
        try:
            if tipo == models.TipoImportacao.ESCOLAS:
                codigo = str(row["CODIGO"]).strip(); nome = str(row["NOME"]).strip()
                escola = (await db.execute(select(models.Escola).where(models.Escola.codigo == codigo))).scalar_one_or_none()
                if escola: escola.nome=nome; escola.ativo=True
                else: db.add(models.Escola(nome=nome, codigo=codigo))
            elif tipo == models.TipoImportacao.TURMAS:
                codigo=str(row["ESCOLA_CODIGO"]).strip(); nome=str(row["NOME"]).strip(); ano=parse_int(row["ANO_LETIVO"],"ANO_LETIVO")
                escola=(await db.execute(select(models.Escola).where(models.Escola.codigo==codigo))).scalar_one_or_none()
                if not escola: raise ValueError(f"escola não encontrada: {codigo}")
                turma=(await db.execute(select(models.Turma).where(models.Turma.escola_id==escola.id, models.Turma.nome==nome, models.Turma.ano_letivo==ano))).scalar_one_or_none()
                if turma: turma.turno=str(row.get("TURNO") or "").strip() or None; turma.ativo=True
                else: db.add(models.Turma(escola_id=escola.id,nome=nome,ano_letivo=ano,turno=str(row.get("TURNO") or "").strip() or None))
            elif tipo == models.TipoImportacao.ALUNOS:
                codigo=str(row["ESCOLA_CODIGO"]).strip(); matricula=str(row["MATRICULA"]).strip(); nome=str(row["NOME"]).strip()
                escola=(await db.execute(select(models.Escola).where(models.Escola.codigo==codigo))).scalar_one_or_none()
                if not escola: raise ValueError(f"escola não encontrada: {codigo}")
                nascimento=parse_data(row.get("DATA_NASCIMENTO")) if row.get("DATA_NASCIMENTO") not in (None,"") else None
                aluno=(await db.execute(select(models.Aluno).where(models.Aluno.escola_id==escola.id,models.Aluno.matricula==matricula))).scalar_one_or_none()
                if aluno: aluno.nome=nome; aluno.data_nascimento=nascimento; aluno.ativo=True
                else: db.add(models.Aluno(escola_id=escola.id,matricula=matricula,nome=nome,data_nascimento=nascimento))
            elif tipo == models.TipoImportacao.MATRICULAS:
                codigo=str(row["ESCOLA_CODIGO"]).strip(); matricula=str(row["MATRICULA"]).strip(); turma_nome=str(row["TURMA_NOME"]).strip(); ano=parse_int(row["ANO_LETIVO"],"ANO_LETIVO")
                escola=(await db.execute(select(models.Escola).where(models.Escola.codigo==codigo))).scalar_one_or_none()
                if not escola: raise ValueError(f"escola não encontrada: {codigo}")
                aluno=(await db.execute(select(models.Aluno).where(models.Aluno.escola_id==escola.id,models.Aluno.matricula==matricula))).scalar_one_or_none()
                if not aluno: raise ValueError(f"aluno não encontrado: {matricula}")
                turma=(await db.execute(select(models.Turma).where(models.Turma.escola_id==escola.id,models.Turma.nome==turma_nome,models.Turma.ano_letivo==ano))).scalar_one_or_none()
                if not turma: raise ValueError(f"turma não encontrada: {turma_nome}/{ano}")
                inicio=parse_data(row.get("DATA_INICIO")) if row.get("DATA_INICIO") not in (None,"") else date(ano,1,1)
                fim=parse_data(row.get("DATA_FIM")) if row.get("DATA_FIM") not in (None,"") else None
                m=(await db.execute(select(models.Matricula).where(models.Matricula.aluno_id==aluno.id,models.Matricula.turma_id==turma.id,models.Matricula.ano_letivo==ano))).scalar_one_or_none()
                if m: m.data_inicio=inicio; m.data_fim=fim; m.status=models.StatusMatricula.ATIVA
                else: db.add(models.Matricula(aluno_id=aluno.id,turma_id=turma.id,ano_letivo=ano,data_inicio=inicio,data_fim=fim,status=models.StatusMatricula.ATIVA))
            importados += 1
        except Exception as exc:  # rollback é responsabilidade do endpoint
            if len(erros) < MAX_ERRORS_RETURNED:
                erros.append({"linha": idx, "erro": str(exc)})
    return importados, erros
